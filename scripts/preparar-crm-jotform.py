"""Lê exports XLSX sem alterar fontes. Requer openpyxl. Gera JSON privado para importador."""
import argparse, hashlib, json, pathlib, datetime, re
import openpyxl
p=argparse.ArgumentParser();p.add_argument('manifesto');p.add_argument('saida');a=p.parse_args()
manifesto=json.loads(pathlib.Path(a.manifesto).read_text());fontes=[];registros=[]
def valor(v):
    if isinstance(v,(datetime.datetime,datetime.date)):return v.isoformat()
    return v
for f in manifesto:
    path=pathlib.Path(f['arquivo']);sha=hashlib.sha256(path.read_bytes()).hexdigest()
    w=openpyxl.load_workbook(path,read_only=False,data_only=True); wf=openpyxl.load_workbook(path,read_only=False,data_only=False)
    sheet=w.active; fs=wf.active; rows=list(sheet.rows);headers=[str(c.value or '') for c in rows[0]]
    if len(rows)-1!=f['esperado']:raise ValueError('Contagem diferente: '+f['nome'])
    if headers.count('CNPJ')!=1:raise ValueError('CNPJ ausente ou repetido')
    for linha,row in enumerate(rows[1:],2):
        campos=[]
        for i,c in enumerate(row):
            orig=valor(c.value);fc=fs.cell(linha,i+1);nome=headers[i];key=nome.strip().upper()
            campos.append({'nome':nome,'valor':orig,'original':orig,'tipo':c.data_type,'formato':c.number_format,
                'formula':fc.value if fc.data_type=='f' else None,'link':fc.hyperlink.target if fc.hyperlink else None,
                'protegido':key in ['N°','Nº','CNPJ','EMPRESA','EMPRESAS','COLABORADOR'] or fc.data_type=='f'})
        cnpj=re.sub(r'\D','',str(row[headers.index('CNPJ')].value or ''))
        if isinstance(row[headers.index('CNPJ')].value,int):cnpj=cnpj.zfill(14)
        def get(*nomes):
            return next((campos[i]['valor'] for i,h in enumerate(headers) if h.strip() in nomes),None)
        registros.append({'id':f['id']+'_'+str(linha).zfill(6),'fonte':f['id'],'fonte_nome':f['nome'],'origem_sha256':sha,
            'linha':linha,'cnpj':cnpj,'nome':str(get('EMPRESA','EMPRESAS') or ''),'codigo':str(get('N°','Nº') or ''),
            'campos':campos,'revisao':1})
    fontes.append({'id':f['id'],'nome':f['nome'],'sha256':sha,'registros':len(rows)-1,'colunas':len(headers),'cabecalhos':headers})
out=pathlib.Path(a.saida);out.write_text(json.dumps({'fontes':fontes,'registros':registros},ensure_ascii=False,indent=2));out.chmod(0o600)
print(json.dumps({'fontes':len(fontes),'registros':len(registros),'cnpjs_distintos':len(set(r['cnpj'] for r in registros))}))
